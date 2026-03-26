from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

import rag_part1
import rag_part2

def create_document():
    doc = Document()

    # Title
    title = doc.add_heading('Building a RAG-Based Solution for Technical Manuals', 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    # Task 1
    doc.add_heading('Task 1 – Diagnose the Keyword Search Failure', level=1)

    t1_results = rag_part1.task1_keyword_vs_semantic()
    doc.add_paragraph(f"Query: {t1_results['query']}")
    doc.add_paragraph(f"Document: {t1_results['document']}")
    doc.add_paragraph(f"Keyword Search Score (Overlap): {t1_results['keyword_score']} (Overlap tokens: {t1_results['keyword_overlap']})")
    doc.add_paragraph(f"Semantic Search Score (Cosine Similarity): {t1_results['semantic_score']:.4f}")

    doc.add_heading('Explanation:', level=2)
    doc.add_paragraph(t1_results['explanation'])

    # Task 2
    doc.add_heading('Task 2 – Build the Chunking Pipeline', level=1)

    t2_results = rag_part1.task2_chunking()
    doc.add_paragraph("Original Paragraph:")
    doc.add_paragraph(t2_results['paragraph'])

    doc.add_heading('Fixed-size Chunking (100 chars, 20 overlap):', level=2)
    for i, chunk in enumerate(t2_results['fixed_chunks']):
        doc.add_paragraph(f"[{i}] {chunk}", style='List Bullet')

    doc.add_heading('Sentence-aware Chunking (2 sentences, 1 overlap):', level=2)
    for i, chunk in enumerate(t2_results['sentence_chunks']):
        doc.add_paragraph(f"[{i}] {chunk}", style='List Bullet')

    doc.add_heading('Explanation:', level=2)
    doc.add_paragraph(t2_results['explanation'])

    # Task 3
    doc.add_heading('Task 3 – Build and Query the Vector Store', level=1)

    t3_results = rag_part2.task3_vector_store()
    doc.add_paragraph(f"Query: {t3_results['query']}")

    doc.add_heading('Top 3 Retrieved Chunks:', level=2)
    for i, res in enumerate(t3_results['results']):
        p = doc.add_paragraph(f"Rank {i+1} (Score: {res['score']:.4f}): {res['chunk']}")

    doc.add_heading('Explanation:', level=2)
    doc.add_paragraph(t3_results['explanation'])

    # Task 4
    doc.add_heading('Task 4 – End-to-End Pipeline with Grounded Generation', level=1)

    t4_results = rag_part2.task4_end_to_end_rag()

    doc.add_heading('Prompt:', level=2)
    prompt_p = doc.add_paragraph(t4_results['prompt'])
    prompt_p.style.font.name = 'Courier New'
    prompt_p.style.font.size = Pt(9)

    doc.add_heading('LLM Output:', level=2)
    doc.add_paragraph(t4_results['llm_response'])

    doc.add_heading('Pipeline Diagram:', level=2)
    rag_part2.draw_diagram()
    doc.add_picture('rag_diagram.png', width=Inches(6.0))

    # Task 5
    doc.add_heading('Task 5 – Reflection', level=1)

    reflection1 = (
        "1. Direct LLM vs. RAG-based Query: A direct LLM query relies entirely on the model's internal, pre-trained parameters (its memory). "
        "For proprietary, highly specific, or updated technical content, a direct LLM query will likely hallucinate or fail because the "
        "information simply isn't in its training set. A RAG (Retrieval-Augmented Generation) pipeline changes this by using the LLM "
        "not as a database, but as a reasoning and synthesizing engine. It actively retrieves the exact technical manual sections needed "
        "and injects them into the prompt as context, grounding the LLM's answer in explicit reality."
    )

    reflection2 = (
        "2. Confidence vs. Correctness: LLMs are designed to generate text that is statistically probable and coherent, which reads as 'confident'. "
        "However, confidence is just a measure of token probability, not factual truth. An LLM might confidently string together plausible-sounding "
        "technical jargon (e.g., advising a user to 'calibrate the acoustic bearings') because those words often appear together in its training data. "
        "In a support environment, a highly confident, factually wrong answer is much more dangerous than a correct but hesitant one, as it can lead "
        "customers to break their equipment or void warranties."
    )

    reflection3 = (
        "3. The Need for Human Review: Even in a well-built RAG system, human review remains critical. RAG depends entirely on the quality "
        "of the retrieved chunks. If a manual contains contradictory information across different chapters, or if an outdated version of a document "
        "is still in the vector database, the LLM will generate a grounded, perfectly reasoned, but ultimately incorrect answer based on bad context. "
        "Furthermore, edge cases or high-stakes physical safety questions often require human judgment to ensure the nuanced context of the "
        "customer's physical environment is taken into account before taking action."
    )

    doc.add_paragraph(reflection1)
    doc.add_paragraph(reflection2)
    doc.add_paragraph(reflection3)

    doc.save('RAG_Solution.docx')
    print("Document successfully saved as RAG_Solution.docx")

if __name__ == "__main__":
    create_document()
